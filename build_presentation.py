"""Build the Gujarati faculty presentation from the DriveAware project README."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUTPUT = Path(__file__).with_name("Driver_Drowsiness_Gujarati_Presentation.pptx")
FONT = "Nirmala UI"
BG = RGBColor(7, 16, 14)
PANEL = RGBColor(13, 27, 22)
PANEL2 = RGBColor(18, 37, 29)
GREEN = RGBColor(104, 221, 157)
WHITE = RGBColor(235, 243, 238)
MUTED = RGBColor(158, 179, 166)
YELLOW = RGBColor(245, 194, 91)
RED = RGBColor(255, 105, 111)
LINE = RGBColor(38, 62, 49)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def rect(slide, x, y, w, h, color, line_color=None, rounded=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line_color or color
    if rounded:
        shape.adjustments[0] = 0.12
    return shape


def text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=18,
    color=WHITE,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    paragraph.space_before = Pt(0)
    paragraph.text = value
    for run in paragraph.runs:
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def base(title, section, number, subtitle=None):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    rect(slide, 0.5, 0.48, 0.11, 0.34, GREEN, rounded=False)
    text(slide, "DRIVEAWARE  /  પ્રોજેક્ટ પ્રેઝન્ટેશન", 0.73, 0.45, 5.8, 0.31, 10, GREEN, True)
    text(slide, section.upper(), 10.25, 0.45, 2.55, 0.31, 9, MUTED, align=PP_ALIGN.RIGHT)
    text(slide, title, 0.65, 1.02, 12.0, 0.62, 28, WHITE, True)
    if subtitle:
        text(slide, subtitle, 0.68, 1.68, 12.0, 0.42, 13, MUTED)
    rect(slide, 0.65, 7.08, 12.03, 0.012, LINE, rounded=False)
    text(
        slide,
        "શૈક્ષણિક પ્રોટોટાઇપ  •  સલામતી માટેનું એકમાત્ર સાધન નથી",
        0.68,
        7.14,
        10.3,
        0.2,
        8,
        MUTED,
    )
    text(slide, f"{number:02d}", 12.08, 7.11, 0.56, 0.24, 9, GREEN, True, PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, heading, body, accent=GREEN, body_size=15):
    rect(slide, x, y, w, h, PANEL, line_color=LINE)
    rect(slide, x + 0.22, y + 0.24, 0.07, 0.35, accent, rounded=False)
    text(slide, heading, x + 0.43, y + 0.2, w - 0.65, 0.48, 17, WHITE, True)
    text(slide, body, x + 0.25, y + 0.82, w - 0.5, h - 0.99, body_size, MUTED)


# 1. Cover
slide = prs.slides.add_slide(BLANK)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG
rect(slide, 8.5, 0, 4.84, 7.5, PANEL, rounded=False)
rect(slide, 0.78, 0.82, 0.13, 0.44, GREEN, rounded=False)
text(slide, "AI + COMPUTER VISION", 1.06, 0.84, 4.4, 0.35, 11, GREEN, True)
text(slide, "ડ્રાઇવર ઊંઘ\nઓળખ પ્રણાલી", 0.84, 1.75, 7.2, 1.68, 38, WHITE, True)
text(
    slide,
    "રીઅલ-ટાઇમ ડ્રાઇવર મોનિટરિંગ માટેનો\nવેબ-આધારિત શૈક્ષણિક પ્રોટોટાઇપ",
    0.88,
    3.75,
    6.8,
    0.98,
    21,
    MUTED,
)
rect(slide, 0.88, 5.22, 4.0, 0.54, PANEL2, line_color=LINE)
text(slide, "Python  •  Flask  •  MediaPipe  •  JavaScript", 1.05, 5.34, 3.7, 0.3, 11, GREEN, True)
text(slide, "પ્રસ્તુતકર્તા: ____________________", 0.9, 6.63, 5.3, 0.3, 12, MUTED)
for diameter, color in [
    (2.25, LINE),
    (1.72, RGBColor(31, 66, 49)),
    (1.13, RGBColor(40, 88, 61)),
]:
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(9.42 + (2.25 - diameter) / 2),
        Inches(2.03 + (2.25 - diameter) / 2),
        Inches(diameter),
        Inches(diameter),
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.color.rgb = color
circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.27), Inches(2.88), Inches(0.55), Inches(0.55))
circle.fill.solid()
circle.fill.fore_color.rgb = GREEN
circle.line.color.rgb = GREEN
text(slide, "જાગૃતિ • જોખમ સંકેત • સમયસર ચેતવણી", 8.84, 4.78, 3.7, 0.76, 17, WHITE, True, PP_ALIGN.CENTER)
text(
    slide,
    "સલામતી પ્રોટોટાઇપ — પ્રમાણિત વાહન ઉપકરણ નથી",
    8.82,
    6.74,
    3.8,
    0.28,
    9,
    MUTED,
    align=PP_ALIGN.CENTER,
)

# 2. Problem and goal
slide = base("સમસ્યા અને પ્રોજેક્ટનો હેતુ", "પરિચય", 2, "ડ્રાઇવિંગ દરમિયાન થાકનું વહેલું નિરીક્ષણ કેમ ઉપયોગી છે?")
card(
    slide,
    0.75,
    2.38,
    5.78,
    3.42,
    "સમસ્યા",
    "• થાકથી ધ્યાન અને પ્રતિક્રિયા ધીમી પડી શકે છે.\n\n"
    "• ઊંઘ આવવી અથવા લાંબી આંખ ઝબકાવવી જોખમનો સંકેત બની શકે છે.\n\n"
    "• માત્ર માનવીય સ્વ-અનુભૂતિ પર નિર્ભર રહેવું પૂરતું નથી.",
    RED,
    16,
)
card(
    slide,
    6.8,
    2.38,
    5.78,
    3.42,
    "અમારો હેતુ",
    "• વેબકેમથી ચહેરાના સંકેતોનું સતત નિરીક્ષણ.\n\n"
    "• આંખ બંધ રહેવાનો સમય અને અન્ય થાક સંકેતોનું અંદાજન.\n\n"
    "• જોખમ વધે ત્યારે દૃશ્ય અને અવાજ ચેતવણી.",
    GREEN,
    16,
)
text(
    slide,
    "મુખ્ય વિચાર: સંકેત ઓળખવો → જોખમનું અંદાજન → સમયસર ચેતવણી",
    1.0,
    6.12,
    11.2,
    0.45,
    19,
    GREEN,
    True,
    PP_ALIGN.CENTER,
)

# 3. Project capabilities
slide = base("સિસ્ટમ શું કરે છે?", "મુખ્ય સુવિધાઓ", 3, "લાઇવ કેમેરાથી ડેશબોર્ડ સુધીની મુખ્ય ક્ષમતાઓ")
card(slide, 0.75, 2.3, 5.78, 1.72, "ચહેરો અને આંખો", "Face Landmarker વડે ચહેરાના પોઇન્ટ; આંખ ખુલ્લી/બંધ અને EAR ગણતરી.", GREEN, 14)
card(slide, 6.8, 2.3, 5.78, 1.72, "થાકના સંકેતો", "આંખ બંધ સમય, PERCLOS, વારંવાર બગાસાં અને માથાના ઝુકાવનો અંદાજ.", YELLOW, 14)
card(slide, 0.75, 4.3, 5.78, 1.72, "ચેતવણી અને નિયંત્રણ", "દૃશ્ય ચેતવણી, Web Audio એલાર્મ, વૈકલ્પિક વોઇસ અલર્ટ અને ટેસ્ટ બટન.", RED, 14)
card(slide, 6.8, 4.3, 5.78, 1.72, "લાઇવ ડેશબોર્ડ", "ડ્રાઇવર સ્થિતિ, જોખમ સ્કોર, ઇવેન્ટ ઇતિહાસ અને કસ્ટમાઇઝ કરી શકાય તેવી થ્રેશોલ્ડ.", GREEN, 14)

# 4. Architecture and privacy
slide = base(
    "સિસ્ટમ આર્કિટેક્ચર અને ગોપનીયતા",
    "રચના",
    4,
    "વિડિયો પ્રક્રિયા બ્રાઉઝરમાં થાય છે; સર્વરને વિડિયો મોકલાતો નથી.",
)
blocks = [
    (0.78, "વેબકેમ", "લાઇવ ફ્રેમ\nબ્રાઉઝરમાં"),
    (3.18, "MediaPipe", "Face Landmarker\nચહેરાના પોઇન્ટ"),
    (5.58, "JavaScript", "EAR • PERCLOS\nથાકનું અંદાજન"),
    (7.98, "ચેતવણી", "સ્ક્રીન • અવાજ\nવોઇસ અલર્ટ"),
    (10.38, "Flask", "API • સેટિંગ્સ\nઇવેન્ટ મેટાડેટા"),
]
for index, (x, heading, body) in enumerate(blocks):
    rect(slide, x, 3.0, 2.05, 1.55, PANEL2 if index in (1, 2) else PANEL, line_color=LINE)
    text(slide, heading, x + 0.12, 3.23, 1.81, 0.42, 15, GREEN if index != 3 else YELLOW, True, PP_ALIGN.CENTER)
    text(slide, body, x + 0.12, 3.76, 1.81, 0.58, 11, MUTED, align=PP_ALIGN.CENTER)
    if index < len(blocks) - 1:
        text(slide, "→", x + 2.06, 3.48, 0.34, 0.4, 20, GREEN, True, PP_ALIGN.CENTER)
rect(slide, 1.05, 5.2, 11.2, 0.78, PANEL, line_color=LINE)
text(
    slide,
    "ગોપનીયતા: કેમેરા ફ્રેમ સ્થાનિક ઉપકરણ પર જ રહે છે. Flaskને ફક્ત શોધ ઇવેન્ટનું લખાણ/આંકડા મળે છે.",
    1.28,
    5.4,
    10.75,
    0.38,
    15,
    WHITE,
    True,
    PP_ALIGN.CENTER,
)
text(
    slide,
    "નોંધ: MediaPipe લાઇબ્રેરી, WASM અને મોડેલ પ્રથમ ઉપયોગે ઇન્ટરનેટ પરથી લોડ થાય છે.",
    1.25,
    6.24,
    10.8,
    0.35,
    11,
    MUTED,
    align=PP_ALIGN.CENTER,
)

# 5. Eye Aspect Ratio
slide = base(
    "આંખ બંધ ઓળખ: Eye Aspect Ratio",
    "ડિટેક્શન",
    5,
    "છ આંખ-લેન્ડમાર્કથી આંખ ખુલ્લી છે કે બંધ તેનો અંદાજ.",
)
card(
    slide,
    0.78,
    2.34,
    5.8,
    3.62,
    "EAR સૂત્ર",
    "EAR = (ઊભા અંતર 1 + ઊભા અંતર 2)\n                 ÷ (2 × આડું અંતર)\n\n"
    "બંને આંખનો સરેરાશ EAR લેવામાં આવે છે.\n\n"
    "ડિફોલ્ટ થ્રેશોલ્ડ: 0.22\nડિફોલ્ટ સતત બંધ સમય: 1.8 સેકન્ડ",
    GREEN,
    18,
)
card(
    slide,
    6.82,
    2.34,
    5.72,
    3.62,
    "નિર્ણયની પ્રક્રિયા",
    "EAR થ્રેશોલ્ડથી ઉપર → આંખ ખુલ્લી.\n\n"
    "EAR થ્રેશોલ્ડથી નીચે → બંધ સમય માપવાનું શરૂ.\n\n"
    "સમય સેટ કરેલી મર્યાદા વટાવે → જોખમ સ્થિતિ અને એલાર્મ.\n\n"
    "થ્રેશોલ્ડ ડેશબોર્ડમાંથી બદલી શકાય છે.",
    YELLOW,
    16,
)
text(
    slide,
    "ટૂંકા પલકારાને બદલે સતત આંખ બંધ રહેવાનો સમય ધ્યાનમાં લેવાય છે.",
    1.0,
    6.25,
    11.3,
    0.4,
    17,
    GREEN,
    True,
    PP_ALIGN.CENTER,
)

# 6. Other indicators
slide = base(
    "અન્ય થાક સંકેતો અને જોખમ સ્કોર",
    "ડિટેક્શન",
    6,
    "બહુવિધ સંકેતોનું સંયોજન — નિદાન નહીં, માત્ર હ્યુરિસ્ટિક અંદાજ.",
)
card(slide, 0.75, 2.28, 5.78, 1.75, "PERCLOS", "રોલિંગ 60 સેકન્ડમાં ચહેરો દેખાતો હોય ત્યારે બંધ આંખના નમૂનાઓનું પ્રમાણ.", GREEN, 14)
card(slide, 6.8, 2.28, 5.78, 1.75, "બગાસાં", "મોઢું કેટલું ખુલ્યું તે માપે; 5 મિનિટમાં વારંવાર યૉન જોખમ સ્કોર વધારે.", YELLOW, 14)
card(slide, 0.75, 4.32, 5.78, 1.75, "માથાની સ્થિતિ", "આંખની રેખાનો ઝુકાવ અને નાકની સાપેક્ષ સ્થિતિમાં ફેરફારનો અંદાજ.", YELLOW, 14)
card(slide, 6.8, 4.32, 5.78, 1.75, "જોખમ સ્કોર", "આંખ બંધ સમય + PERCLOS + બગાસાં + માથાની હિલચાલનું હ્યુરિસ્ટિક સંયોજન.", RED, 14)
text(
    slide,
    "જોખમ ટકાવારી વાસ્તવિક સંભાવના, તબીબી માપ કે પ્રમાણિત માપદંડ નથી.",
    1.0,
    6.42,
    11.3,
    0.35,
    13,
    RED,
    True,
    PP_ALIGN.CENTER,
)

# 7. Workflow and alert
slide = base("કાર્યપ્રવાહ અને ચેતવણી", "સિસ્ટમ વર્તન", 7, "મોનિટરિંગ શરૂ થયા પછી શું થાય છે?")
steps = [
    ("01", "કેમેરા ચાલુ", "વપરાશકર્તાની મંજૂરી"),
    ("02", "ચહેરો શોધો", "એક ચહેરો જરૂરી"),
    ("03", "સંકેત માપો", "આંખ • મોઢું • માથું"),
    ("04", "જોખમ આંકો", "થ્રેશોલ્ડ તપાસો"),
]
for index, (number, heading, body) in enumerate(steps):
    x = 0.86 + index * 3.12
    rect(slide, x, 2.6, 2.65, 1.52, PANEL2, line_color=LINE)
    text(slide, number, x + 0.18, 2.78, 0.5, 0.35, 12, GREEN, True)
    text(slide, heading, x + 0.18, 3.13, 2.28, 0.38, 16, WHITE, True)
    text(slide, body, x + 0.18, 3.57, 2.28, 0.31, 11, MUTED)
    if index < 3:
        text(slide, "→", x + 2.72, 3.13, 0.34, 0.4, 20, GREEN, True, PP_ALIGN.CENTER)
rect(slide, 1.1, 4.75, 11.15, 0.98, RGBColor(43, 21, 22), line_color=RGBColor(105, 49, 50))
text(
    slide,
    "જોખમ વધે ત્યારે: “DROWSINESS DETECTED!”",
    1.4,
    4.95,
    10.55,
    0.34,
    19,
    RED,
    True,
    PP_ALIGN.CENTER,
)
text(
    slide,
    "સ્ક્રીન પર ચેતવણી  •  વારંવાર વાગતું અવાજ એલાર્મ  •  વૈકલ્પિક વોઇસ સૂચના",
    1.3,
    5.98,
    10.75,
    0.37,
    14,
    MUTED,
    align=PP_ALIGN.CENTER,
)

# 8. Dashboard and settings
slide = base("ડેશબોર્ડ, નિયંત્રણ અને ગોઠવણ", "વપરાશકર્તા અનુભવ", 8, "એક સ્ક્રીન પર લાઇવ સ્થિતિ અને પ્રોટોટાઇપ નિયંત્રણો.")
card(
    slide,
    0.78,
    2.32,
    5.78,
    3.7,
    "ડેશબોર્ડ પર",
    "• લાઇવ કેમેરા અને લેન્ડમાર્ક\n\n"
    "• ચહેરો/કેમેરા સ્થિતિ\n\n"
    "• આંખ, EAR, આંખ બંધ સમય\n\n"
    "• જોખમ, બગાસાં, માથાની સ્થિતિ, PERCLOS\n\n"
    "• ઇવેન્ટ ઇતિહાસ",
    GREEN,
    15,
)
card(
    slide,
    6.8,
    2.32,
    5.78,
    3.7,
    "નિયંત્રણ અને થ્રેશોલ્ડ",
    "• કેમેરા/મોનિટરિંગ શરૂ અથવા બંધ\n\n"
    "• એલાર્મ ટેસ્ટ; અવાજ અને વોઇસ ચાલુ/બંધ\n\n"
    "• આંકડા રીસેટ\n\n"
    "• EAR, બંધ સમય, PERCLOS વિન્ડો, જોખમ મર્યાદા\n\n"
    "• યૉન મર્યાદા, એલાર્મ કૂલડાઉન, હેડ ટિલ્ટ",
    YELLOW,
    14,
)
text(
    slide,
    "સેટિંગ અને ઇવેન્ટ ઇતિહાસ સર્વર મેમરીમાં રહે છે; સર્વર બંધ થતાં રીસેટ થાય છે.",
    1.0,
    6.32,
    11.3,
    0.36,
    13,
    MUTED,
    align=PP_ALIGN.CENTER,
)

# 9. Technology and demo
slide = base("ટેકનોલોજી અને ડેમો ચલાવવાની રીત", "અમલીકરણ", 9, "આ પ્રોટોટાઇપ Windows પર સરળતાથી ચલાવવા માટે બનાવ્યો છે.")
card(
    slide,
    0.78,
    2.25,
    5.78,
    3.45,
    "ટેક સ્ટેક",
    "Python + Flask — સર્વર અને APIs\n\n"
    "HTML / CSS / JavaScript — ડેશબોર્ડ\n\n"
    "MediaPipe Face Landmarker — ચહેરા લેન્ડમાર્ક\n\n"
    "Web Audio API — બ્રાઉઝર એલાર્મ\n\n"
    "ઇવેન્ટ્સ/સેટિંગ્સ — ઇન-મેમરી",
    GREEN,
    15,
)
rect(slide, 6.8, 2.25, 5.78, 3.45, PANEL, line_color=LINE)
text(slide, "ડેમો માટેના પગલાં", 7.05, 2.5, 5.2, 0.45, 17, WHITE, True)
text(
    slide,
    "1. Python 3.11 ઇન્સ્ટોલ કરો.\n\n"
    "2. requirements.txt માંથી Flask ઇન્સ્ટોલ કરો.\n\n"
    "3. python app.py ચલાવો.\n\n"
    "4. localhost:5000 ખોલી કેમેરા પરમિશન આપો.\n\n"
    "5. Start monitoring દબાવો.",
    7.05,
    3.06,
    5.15,
    2.3,
    14,
    MUTED,
)
rect(slide, 0.95, 5.98, 11.45, 0.67, PANEL2, line_color=LINE)
text(
    slide,
    "પ્રથમ વખત MediaPipe JavaScript/WASM અને મોડેલ લોડ કરવા ઇન્ટરનેટ જરૂરી છે.",
    1.15,
    6.15,
    11.05,
    0.33,
    13,
    GREEN,
    True,
    PP_ALIGN.CENTER,
)

# 10. Limitations and closing
slide = base(
    "મર્યાદાઓ, ભવિષ્ય અને નિષ્કર્ષ",
    "સમાપન",
    10,
    "સલામતી દાવા કરતાં જવાબદાર પ્રોટોટાઇપ અને યોગ્ય માન્યતા વધુ મહત્વની છે.",
)
card(
    slide,
    0.75,
    2.24,
    5.78,
    2.28,
    "હાલની મર્યાદાઓ",
    "• પ્રકાશ, ચશ્માં, કેમેરા કોણ અને વ્યક્તિગત તફાવત અસર કરે છે.\n\n"
    "• એકથી વધુ ચહેરા હોય તો માપન રોકાય છે.\n\n"
    "• હ્યુરિસ્ટિક સ્કોર વાસ્તવિક જીવનમાં માન્ય કરેલો નથી.",
    RED,
    13,
)
card(
    slide,
    6.8,
    2.24,
    5.78,
    2.28,
    "ભવિષ્યમાં",
    "• વિવિધ પ્રકાશ/ચશ્માં/કેમેરા પર પરીક્ષણ.\n\n"
    "• વાસ્તવિક ડેટા સાથે થ્રેશોલ્ડનું કેલિબ્રેશન.\n\n"
    "• સુલભતા, ભૂલ સંભાળ અને કામગીરી સુધારવી.",
    GREEN,
    13,
)
rect(slide, 1.0, 4.9, 11.35, 1.12, RGBColor(43, 21, 22), line_color=RGBColor(105, 49, 50))
text(slide, "મહત્વની સલામતી નોંધ", 1.28, 5.07, 10.8, 0.34, 16, RED, True, PP_ALIGN.CENTER)
text(
    slide,
    "આ પ્રમાણિત ઓટોમોટિવ સલામતી ઉપકરણ નથી. ડ્રાઇવિંગમાં થાક સામે એકમાત્ર સુરક્ષા તરીકે તેનો ઉપયોગ ન કરો.",
    1.3,
    5.48,
    10.7,
    0.39,
    13,
    WHITE,
    align=PP_ALIGN.CENTER,
)
text(slide, "આભાર  •  પ્રશ્નો?", 1.0, 6.4, 11.3, 0.43, 22, GREEN, True, PP_ALIGN.CENTER)

prs.core_properties.title = "ડ્રાઇવર ઊંઘ ઓળખ પ્રણાલી — ગુજરાતી પ્રેઝન્ટેશન"
prs.core_properties.subject = "Flask અને MediaPipe આધારિત ડ્રાઇવર મોનિટરિંગ શૈક્ષણિક પ્રોટોટાઇપ"
prs.core_properties.author = "DriveAware Project"
prs.core_properties.keywords = "Gujarati, Driver Drowsiness, MediaPipe, Flask, Presentation"
prs.save(OUTPUT)
print(f"Created {OUTPUT} ({len(prs.slides)} slides, {OUTPUT.stat().st_size:,} bytes).")
